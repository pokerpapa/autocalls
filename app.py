# app.py
import os
import uuid
import json
import tempfile
import logging
from datetime import datetime
from flask import (
    Flask, request, render_template, jsonify,
    send_from_directory, url_for, abort
)
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from docx import Document

from vostools import mp3_to_wav, transcribe
from mistral_tools import call_mistral

load_dotenv()
logging.basicConfig(level=logging.DEBUG)

app = Flask(__name__)

# Список менеджеров — можно править под ваши реальные имена
MANAGERS = [
    "Иван Иванов",
    "Мария Петрова",
    "Алексей Смирнов",
    "Анна Кузнецова"
]

UPLOAD_DIR  = os.path.join(app.instance_path, "uploads")
RESULTS_DIR = os.path.join(app.instance_path, "results")
os.makedirs(UPLOAD_DIR,  exist_ok=True)
os.makedirs(RESULTS_DIR, exist_ok=True)


@app.get("/")
def index():
    return render_template("upload.html", managers=MANAGERS)


@app.post("/api/process")
def api_process():
    """
    Принимаем один MP3 и поле manager, генерируем отчёт,
    сохраняем .docx в папке results/<manager> и возвращаем ссылку на скачивание.
    """
    audio   = request.files.get("audio")
    manager = request.form.get("manager", "").strip()
    if not audio or not manager:
        return jsonify({"error": "Нужно указать менеджера и файл"}), 400

    # Защищённое имя менеджера для папки
    mgr_slug = secure_filename(manager)
    mgr_dir  = os.path.join(RESULTS_DIR, mgr_slug)
    os.makedirs(mgr_dir, exist_ok=True)

    # Сохраняем MP3 под оригинальным именем + timestamp
    orig_name = secure_filename(audio.filename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    mp3_name  = f"{timestamp}_{orig_name}"
    mp3_path  = os.path.join(UPLOAD_DIR, mp3_name)
    audio.save(mp3_path)

    try:
        # Конверт и ASR
        wav_file = tempfile.mktemp(suffix=".wav")
        mp3_to_wav(mp3_path, wav_file)
        asr_json = transcribe(wav_file)
        raw_text = json.loads(asr_json).get("text", "")

        # Запрос в Mistral
        analysis = call_mistral(raw_text)

        # Собираем Word‑документ
        doc = Document()
        doc.add_heading(f"Отчёт по звонку — {manager}", level=1)
        doc.add_paragraph(analysis)

        base     = os.path.splitext(mp3_name)[0]
        doc_name = f"{base}.docx"
        doc_path = os.path.join(mgr_dir, doc_name)
        doc.save(doc_path)

        # Публичная ссылка
        download_url = url_for(
            'download_result',
            manager=mgr_slug,
            filename=doc_name,
            _external=True
        )
        return jsonify({"download_url": download_url})

    except Exception as e:
        logging.exception("Ошибка при обработке")
        return jsonify({"error": str(e)}), 500

    finally:
        # Чистим временные файлы
        for fp in (mp3_path, wav_file):
            try: os.remove(fp)
            except: pass


@app.get("/results/<manager>/<path:filename>")
def download_result(manager, filename):
    """
    Отдаёт .docx из папки results/<manager>
    """
    safe_mgr = secure_filename(manager)
    return send_from_directory(
        os.path.join(RESULTS_DIR, safe_mgr),
        filename,
        as_attachment=True
    )


@app.get("/jobs")
def list_jobs():
    """
    Страница со списком всех отчётов, сгруппированных по менеджерам.
    """
    all_data = []
    for manager in MANAGERS:
        mgr_slug = secure_filename(manager)
        mgr_dir  = os.path.join(RESULTS_DIR, mgr_slug)
        entries  = []
        if os.path.isdir(mgr_dir):
            for fname in sorted(os.listdir(mgr_dir), reverse=True):
                if fname.lower().endswith(".docx"):
                    path = os.path.join(mgr_dir, fname)
                    mtime = datetime.fromtimestamp(os.path.getmtime(path))
                    entries.append({
                        "name": fname,
                        "url":  url_for('download_result', manager=mgr_slug, filename=fname),
                        "date": mtime.strftime("%Y-%m-%d %H:%M:%S")
                    })
        all_data.append({
            "manager": manager,
            "files": entries
        })
    return render_template("jobs.html", data=all_data)


if __name__ == "__main__":
    app.run(debug=True, port=8000)
